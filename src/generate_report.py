from __future__ import annotations

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.portfolio_construction import class_allocation, portfolio_returns, weights_to_frame
from src.utils import ASSET_CLASS_MAP, ASSET_CLASS_ORDER, format_weight_dict


PERCENT_COLUMNS = {
    "weight",
    "current_weight",
    "target_weight",
    "portfolio_volatility",
    "marginal_risk_contribution",
    "total_risk_contribution",
    "percent_risk_contribution",
    "rolling_12m_return",
    "rolling_12m_volatility",
    "rolling_max_drawdown",
    "actual_value",
    "drift",
    "post_trade_weight",
    "annualized_return",
    "annualized_volatility",
    "max_drawdown",
    "tracking_error_vs_benchmark",
    "historical_var_95",
    "historical_cvar_95",
    "best_month",
    "worst_month",
    "positive_month_percentage",
    "correlation_to_benchmark",
    "scenario_return",
    "portfolio_weight",
    "benchmark_weight",
    "etf_contribution",
    "benchmark_etf_contribution",
    "asset_class_contribution",
    "portfolio_scenario_return",
    "benchmark_scenario_return",
    "relative_performance",
    "max_volatility_target",
    "max_drawdown_tolerance",
    "min_equity_allocation",
    "max_equity_allocation",
    "min_fixed_income_allocation",
    "max_fixed_income_allocation",
    "min_alternatives_allocation",
    "max_alternatives_allocation",
    "min_cash_allocation",
    "max_cash_allocation",
    "max_single_etf_weight",
    "rebalancing_threshold",
}

CURRENCY_COLUMNS = {"trade_amount", "estimated_transaction_cost"}


def _set_chart_style() -> None:
    try:
        plt.style.use("seaborn-v0_8-whitegrid")
    except Exception:
        plt.style.use("default")


def _drawdown_series(returns: pd.Series) -> pd.Series:
    wealth = (1.0 + returns).cumprod()
    return wealth / wealth.cummax() - 1.0


def generate_charts(
    charts_dir: Path,
    profile_name: str,
    returns: pd.DataFrame,
    primary_candidates: dict[str, pd.Series],
    primary_benchmark: pd.Series,
    risk_summary: pd.DataFrame,
    factor_summary: pd.DataFrame,
    equity_factor_summary: pd.DataFrame,
    risk_contribution: pd.DataFrame,
    rolling_risk_metrics: pd.DataFrame,
    stress_results: pd.DataFrame,
    rebalancing_trades: pd.DataFrame,
    efficient_frontier: pd.DataFrame,
) -> dict[str, Path]:
    charts_dir.mkdir(parents=True, exist_ok=True)
    _set_chart_style()
    chart_paths: dict[str, Path] = {}
    chart_caption = f"{returns.attrs.get('provenance', {}).get('mode', 'unverified').upper()} input data | same-sample diagnostics"

    path = charts_dir / "efficient_frontier.png"
    fig, ax = plt.subplots(figsize=(9, 6))
    if not efficient_frontier.empty:
        scatter = ax.scatter(
            efficient_frontier["annualized_volatility"],
            efficient_frontier["annualized_return"],
            c=efficient_frontier["sharpe_ratio"],
            cmap="viridis",
            alpha=0.65,
            s=18,
            label="Feasible portfolios",
        )
        fig.colorbar(scatter, ax=ax, label="Sharpe Ratio")
    primary_risk = risk_summary[risk_summary["profile_name"] == profile_name]
    for _, row in primary_risk.iterrows():
        ax.scatter(row["annualized_volatility"], row["annualized_return"], s=80, marker="D")
        ax.annotate(
            row["portfolio_name"].replace(" Portfolio", ""),
            (row["annualized_volatility"], row["annualized_return"]),
            xytext=(6, 5),
            textcoords="offset points",
            fontsize=8,
        )
    ax.set_title(f"Efficient Frontier - {profile_name}")
    ax.set_xlabel("Annualized Volatility")
    ax.set_ylabel("Annualized Return")
    ax.xaxis.set_major_formatter(lambda x, _: f"{x:.0%}")
    ax.yaxis.set_major_formatter(lambda y, _: f"{y:.0%}")
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["efficient_frontier"] = path

    recommended = primary_candidates["IPS Recommended Portfolio"].sort_values()
    path = charts_dir / "recommended_allocation.png"
    fig, ax = plt.subplots(figsize=(9, 6))
    colors = recommended.index.map(
        lambda t: {"Equity": "#1f77b4", "Fixed Income": "#2ca02c", "Alternatives": "#9467bd", "Cash": "#7f7f7f"}[
            ASSET_CLASS_MAP[t]
        ]
    )
    ax.barh(recommended.index, recommended.values, color=list(colors))
    ax.set_title(f"Recommended Allocation - {profile_name}")
    ax.set_xlabel("Portfolio Weight")
    ax.xaxis.set_major_formatter(lambda x, _: f"{x:.0%}")
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["recommended_allocation"] = path

    recommended_returns = portfolio_returns(returns, primary_candidates["IPS Recommended Portfolio"])
    aligned = pd.concat(
        [recommended_returns.rename("Recommended"), primary_benchmark.rename("Benchmark")], axis=1
    ).dropna()
    path = charts_dir / "cumulative_performance_vs_benchmark.png"
    fig, ax = plt.subplots(figsize=(10, 6))
    cumulative = (1.0 + aligned).cumprod()
    cumulative.plot(ax=ax, linewidth=2)
    ax.set_title(f"Cumulative Performance vs Benchmark - {profile_name}")
    ax.set_ylabel("Growth of $1")
    ax.set_xlabel("")
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["cumulative_performance"] = path

    path = charts_dir / "drawdown_comparison.png"
    fig, ax = plt.subplots(figsize=(10, 6))
    pd.DataFrame(
        {
            "Recommended": _drawdown_series(aligned["Recommended"]),
            "Benchmark": _drawdown_series(aligned["Benchmark"]),
        }
    ).plot(ax=ax, linewidth=2)
    ax.set_title(f"Drawdown Comparison - {profile_name}")
    ax.set_ylabel("Drawdown")
    ax.set_xlabel("")
    ax.yaxis.set_major_formatter(lambda y, _: f"{y:.0%}")
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["drawdown_comparison"] = path

    path = charts_dir / "factor_exposures.png"
    fig, ax = plt.subplots(figsize=(8, 5))
    factors = factor_summary[
        (factor_summary["profile_name"] == profile_name) & (factor_summary["factor"] != "Alpha")
    ]
    ax.bar(factors["factor"], factors["coefficient"], color="#4c78a8")
    ax.axhline(0, color="#333333", linewidth=0.8)
    ax.set_title(f"Fama-French Factor Exposures - {profile_name}")
    ax.set_ylabel("Beta")
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["factor_exposures"] = path

    path = charts_dir / "equity_sleeve_factor_exposures.png"
    fig, ax = plt.subplots(figsize=(8, 5))
    equity_factors = equity_factor_summary[
        (equity_factor_summary["profile_name"] == profile_name)
        & (equity_factor_summary["factor"] != "Alpha")
    ]
    ax.bar(equity_factors["factor"], equity_factors["coefficient"], color="#59a14f")
    ax.axhline(0, color="#333333", linewidth=0.8)
    ax.set_title(f"Equity Sleeve Factor Exposures - {profile_name}")
    ax.set_ylabel("Beta")
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["equity_sleeve_factor_exposures"] = path

    path = charts_dir / "risk_contribution.png"
    fig, ax = plt.subplots(figsize=(9, 6))
    risk_chart = risk_contribution.sort_values("percent_risk_contribution")
    colors = risk_chart["asset_class"].map(
        {"Equity": "#1f77b4", "Fixed Income": "#2ca02c", "Alternatives": "#9467bd", "Cash": "#7f7f7f"}
    )
    ax.barh(risk_chart["ticker"], risk_chart["percent_risk_contribution"], color=list(colors))
    ax.set_title(f"Risk Contribution - {profile_name}")
    ax.set_xlabel("Percent Contribution to Portfolio Volatility")
    ax.xaxis.set_major_formatter(lambda x, _: f"{x:.0%}")
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["risk_contribution"] = path

    path = charts_dir / "policy_drift.png"
    fig, ax = plt.subplots(figsize=(9, 5))
    drift = rebalancing_trades[rebalancing_trades["profile_name"] == profile_name].copy()
    drift = drift.sort_values("drift")
    drift_colors = drift["drift_breach_flag"].map({True: "#d62728", False: "#7f7f7f"})
    ax.barh(drift["ticker"], drift["drift"], color=list(drift_colors))
    ax.axvline(0, color="#333333", linewidth=0.8)
    ax.set_title(f"Policy Drift - {profile_name}")
    ax.set_xlabel("Current Weight minus Target Weight")
    ax.xaxis.set_major_formatter(lambda x, _: f"{x:.0%}")
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["policy_drift"] = path

    path = charts_dir / "stress_test_results.png"
    fig, ax = plt.subplots(figsize=(10, 6))
    stress = stress_results[stress_results["profile_name"] == profile_name].drop_duplicates(
        ["scenario_name"]
    )
    stress = stress.sort_values("portfolio_scenario_return")
    ax.barh(stress["scenario_name"], stress["portfolio_scenario_return"], color="#b279a2", label="Recommended")
    ax.scatter(stress["benchmark_scenario_return"], stress["scenario_name"], color="#222222", label="Benchmark")
    ax.axvline(0, color="#333333", linewidth=0.8)
    ax.set_title(f"Stress Test Results - {profile_name}")
    ax.set_xlabel("Scenario Return")
    ax.xaxis.set_major_formatter(lambda x, _: f"{x:.0%}")
    ax.legend()
    fig.suptitle(chart_caption, fontsize=10)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path, dpi=180)
    plt.close(fig)
    chart_paths["stress_test_results"] = path

    rolling = rolling_risk_metrics.copy()
    rolling["date"] = pd.to_datetime(rolling["date"])
    for metric, filename, title, ylabel in [
        ("rolling_12m_return", "rolling_return.png", "Rolling 12-Month Return", "Return"),
        ("rolling_12m_volatility", "rolling_volatility.png", "Rolling 12-Month Volatility", "Volatility"),
        ("rolling_max_drawdown", "rolling_drawdown.png", "Rolling 12-Month Max Drawdown", "Drawdown"),
    ]:
        path = charts_dir / filename
        fig, ax = plt.subplots(figsize=(10, 5))
        for series_name, group in rolling.groupby("series"):
            ax.plot(group["date"], group[metric], label=series_name, linewidth=2)
        ax.axhline(0, color="#333333", linewidth=0.8)
        ax.set_title(f"{title} - {profile_name}")
        ax.set_ylabel(ylabel)
        ax.set_xlabel("")
        ax.yaxis.set_major_formatter(lambda y, _: f"{y:.0%}")
        ax.legend()
        fig.suptitle(chart_caption, fontsize=10)
        fig.tight_layout(rect=[0, 0, 1, 0.95])
        fig.savefig(path, dpi=180)
        plt.close(fig)
        chart_paths[filename.replace(".png", "")] = path

    return chart_paths


def _format_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    negative_fill = PatternFill("solid", fgColor="F4CCCC")
    breach_fill = PatternFill("solid", fgColor="FCE4D6")
    pass_fill = PatternFill("solid", fgColor="D9EAD3")
    fail_fill = PatternFill("solid", fgColor="F4CCCC")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        max_row = sheet.max_row
        max_col = sheet.max_column
        if max_row < 1 or max_col < 1:
            continue
        if sheet.title == "Executive Summary":
            continue
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        sheet.auto_filter.ref = sheet.dimensions

        for column_cells in sheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            header = str(column_cells[0].value or "")
            width = min(max(len(header), *(len(str(cell.value)) for cell in column_cells[1:] if cell.value is not None)) + 2, 48)
            sheet.column_dimensions[column_letter].width = width
            if header in PERCENT_COLUMNS:
                for cell in column_cells[1:]:
                    cell.number_format = "0.00%"
            elif header in CURRENCY_COLUMNS:
                for cell in column_cells[1:]:
                    cell.number_format = '$#,##0;[Red]-$#,##0'
            elif header in {"sharpe_ratio", "sortino_ratio", "information_ratio", "beta_vs_benchmark", "coefficient", "t_stat", "p_value", "r_squared", "rolling_12m_sharpe_ratio"}:
                for cell in column_cells[1:]:
                    cell.number_format = "0.00"

            if header in {
                "annualized_return",
                "best_month",
                "worst_month",
                "historical_var_95",
                "historical_cvar_95",
                "scenario_return",
                "portfolio_scenario_return",
                "benchmark_scenario_return",
                "relative_performance",
            }:
                sheet.conditional_formatting.add(
                    f"{column_letter}2:{column_letter}{max_row}",
                    CellIsRule(operator="lessThan", formula=["0"], fill=negative_fill),
                )
            if header in {"max_drawdown"}:
                sheet.conditional_formatting.add(
                    f"{column_letter}2:{column_letter}{max_row}",
                    CellIsRule(operator="lessThan", formula=["-0.2"], fill=negative_fill),
                )
            if header == "drift_breach_flag":
                sheet.conditional_formatting.add(
                    f"{column_letter}2:{column_letter}{max_row}",
                    CellIsRule(operator="equal", formula=["TRUE"], fill=breach_fill),
                )
            if header == "pass_fail":
                sheet.conditional_formatting.add(
                    f"{column_letter}2:{column_letter}{max_row}",
                    FormulaRule(formula=[f'{column_letter}2="PASS"'], fill=pass_fill),
                )
                sheet.conditional_formatting.add(
                    f"{column_letter}2:{column_letter}{max_row}",
                    FormulaRule(formula=[f'{column_letter}2="FAIL"'], fill=fail_fill),
                )
        if sheet.title in {"Data Provenance", "Source Summary"}:
            sheet.row_dimensions[1].height = 30
            for row in sheet.iter_rows(min_row=2):
                sheet.row_dimensions[row[0].row].height = 60
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(path)


def _build_executive_summary(
    profile_name: str,
    profile: pd.Series,
    recommended_metrics: pd.Series,
    allocation_by_class: dict[str, float],
    stress_results: pd.DataFrame,
    rebalancing_trades: pd.DataFrame,
    constraints_validation: pd.DataFrame,
) -> dict[str, object]:
    benchmark = str(profile["benchmark_definition"]).replace(";", ",")
    equity_bear = stress_results[
        (stress_results["profile_name"] == profile_name)
        & (stress_results["scenario_name"] == "Equity Bear Market")
    ].drop_duplicates("scenario_name")
    equity_bear_lower = False
    if not equity_bear.empty:
        row = equity_bear.iloc[0]
        equity_bear_lower = float(row["portfolio_scenario_return"]) > float(row["benchmark_scenario_return"])
    breach_count = int(
        rebalancing_trades[
            (rebalancing_trades["profile_name"] == profile_name)
            & (rebalancing_trades["drift_breach_flag"])
        ].shape[0]
    )
    constraints_passed = bool((constraints_validation["pass_fail"] == "PASS").all())
    rebalancing_cost = float(
        rebalancing_trades.loc[
            rebalancing_trades["profile_name"] == profile_name,
            "estimated_transaction_cost",
        ].sum()
    )
    return {
        "profile_name": profile_name,
        "portfolio_name": "IPS Recommended Portfolio",
        "portfolio_market_value": 100_000_000.0,
        "benchmark": benchmark,
        "rebalance_threshold": float(profile["rebalancing_threshold"]),
        "kpis": [
            ("Annualized Return", float(recommended_metrics["annualized_return"]), "0.00%"),
            ("Volatility", float(recommended_metrics["annualized_volatility"]), "0.00%"),
            ("Max Drawdown", float(recommended_metrics["max_drawdown"]), "0.00%"),
            ("Tracking Error", float(recommended_metrics["tracking_error_vs_benchmark"]), "0.00%"),
            ("VaR", float(recommended_metrics["historical_var_95"]), "0.00%"),
            ("Rebalancing Cost", rebalancing_cost, '$#,##0'),
        ],
        "allocation": {asset_class: allocation_by_class.get(asset_class, 0.0) for asset_class in ASSET_CLASS_ORDER},
        "committee_notes": [
            "Passed IPS constraint validation" if constraints_passed else "IPS constraint validation requires review",
            "Equity bear market stress loss lower than benchmark"
            if equity_bear_lower
            else "Equity bear market stress result requires review",
            f"{breach_count} drift breaches generated rebalancing trades",
        ],
    }


def _write_executive_summary_dashboard(workbook, summary: dict[str, object]) -> None:
    if "Sheet" in workbook.sheetnames and workbook["Sheet"].max_row == 1 and workbook["Sheet"]["A1"].value is None:
        del workbook["Sheet"]
    if "Executive Summary" in workbook.sheetnames:
        del workbook["Executive Summary"]

    sheet = workbook.create_sheet("Executive Summary", 0)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    for column, width in {"A": 24, "B": 18, "C": 18, "D": 28, "E": 18, "F": 18}.items():
        sheet.column_dimensions[column].width = width

    navy = PatternFill("solid", fgColor="1F4E78")
    blue = PatternFill("solid", fgColor="D9EAF7")
    light = PatternFill("solid", fgColor="F7FBFD")
    red = PatternFill("solid", fgColor="F4CCCC")
    green = PatternFill("solid", fgColor="D9EAD3")
    white_font = Font(bold=True, color="FFFFFF", size=13)
    section_font = Font(bold=True, color="1F4E78")
    label_font = Font(bold=True, color="1F4E78")
    value_font = Font(bold=True, size=14)
    thin_blue = Side(style="thin", color="9EBDD1")
    medium_blue = Side(style="medium", color="1F4E78")
    box_border = Border(left=thin_blue, right=thin_blue, top=medium_blue, bottom=thin_blue)

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Executive Summary - Balanced Growth IPS"
    sheet["A1"].fill = navy
    sheet["A1"].font = white_font
    sheet["A1"].alignment = Alignment(horizontal="center")

    def section(range_ref: str, title: str) -> None:
        sheet.merge_cells(range_ref)
        cell = sheet[range_ref.split(":")[0]]
        cell.value = title
        cell.fill = blue
        cell.font = section_font
        cell.alignment = Alignment(horizontal="left")

    section("A3:C3", "Portfolio Recommendation")
    portfolio_rows = [
        ("IPS Profile", summary["profile_name"]),
        ("Recommended Portfolio", summary["portfolio_name"]),
        ("Portfolio Market Value", summary["portfolio_market_value"]),
        ("Benchmark", summary["benchmark"]),
        ("Rebalance Threshold", summary["rebalance_threshold"]),
    ]
    for row_index, (label, value) in enumerate(portfolio_rows, start=4):
        sheet.cell(row_index, 1, label).font = label_font
        sheet.cell(row_index, 2, value)
        sheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=3)
        sheet.cell(row_index, 2).alignment = Alignment(wrap_text=True)
        if label == "Portfolio Market Value":
            sheet.cell(row_index, 2).number_format = '$#,##0'
        if label == "Rebalance Threshold":
            sheet.cell(row_index, 2).number_format = "0.00%"

    section("D3:F3", "Committee Notes")
    for row_index, note in enumerate(summary["committee_notes"], start=4):
        sheet.cell(row_index, 4, f"- {note}")
        sheet.merge_cells(start_row=row_index, start_column=4, end_row=row_index, end_column=6)
        sheet.cell(row_index, 4).alignment = Alignment(wrap_text=True)
        sheet.cell(row_index, 4).fill = green if "review" not in str(note).lower() else red

    section("A10:F10", "Key Metrics")

    def kpi_box(row: int, col: int, label: str, value: float, number_format: str) -> None:
        sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        sheet.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        label_cell = sheet.cell(row, col, label)
        value_cell = sheet.cell(row + 1, col, value)
        label_cell.font = label_font
        label_cell.fill = blue
        label_cell.alignment = Alignment(horizontal="center")
        value_cell.font = value_font
        value_cell.fill = red if isinstance(value, (int, float)) and value < 0 else light
        value_cell.alignment = Alignment(horizontal="center")
        value_cell.number_format = number_format
        for r in range(row, row + 2):
            for c in range(col, col + 2):
                sheet.cell(r, c).border = box_border

    positions = [(12, 1), (12, 3), (12, 5), (15, 1), (15, 3), (15, 5)]
    for (row, col), (label, value, number_format) in zip(positions, summary["kpis"]):
        kpi_box(row, col, label, value, number_format)

    section("A19:B19", "Asset-Class Allocation")
    sheet.cell(20, 1, "Asset Class").font = label_font
    sheet.cell(20, 2, "Target Weight").font = label_font
    for row_index, asset_class in enumerate(ASSET_CLASS_ORDER, start=21):
        sheet.cell(row_index, 1, asset_class)
        sheet.cell(row_index, 2, summary["allocation"][asset_class])
        sheet.cell(row_index, 2).number_format = "0.00%"
        sheet.cell(row_index, 1).fill = light
        sheet.cell(row_index, 2).fill = light

    for row in sheet.iter_rows(min_row=1, max_row=24, min_col=1, max_col=6):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def generate_excel_report(
    output_path: Path,
    profile_name: str,
    profiles: pd.DataFrame,
    all_weights: dict[str, dict[str, pd.Series]],
    risk_summary: pd.DataFrame,
    factor_summary: pd.DataFrame,
    equity_factor_summary: pd.DataFrame,
    risk_contribution: pd.DataFrame,
    rolling_risk_metrics: pd.DataFrame,
    etf_rationale: pd.DataFrame,
    constraints_validation: pd.DataFrame,
    stress_results: pd.DataFrame,
    rebalancing_trades: pd.DataFrame,
    efficient_frontier: pd.DataFrame,
    chart_paths: dict[str, Path],
    provenance: dict | None = None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    recommended_weights = {
        profile: portfolios["IPS Recommended Portfolio"] for profile, portfolios in all_weights.items()
    }
    allocation = weights_to_frame({k: {"IPS Recommended Portfolio": v} for k, v in recommended_weights.items()})
    recommended_metrics = risk_summary[
        (risk_summary["profile_name"] == profile_name)
        & (risk_summary["portfolio_name"] == "IPS Recommended Portfolio")
    ].iloc[0]
    profile = profiles[profiles["profile_name"] == profile_name].iloc[0]
    allocation_by_class = class_allocation(recommended_weights[profile_name])
    drift = rebalancing_trades[rebalancing_trades["profile_name"] == profile_name]
    summary = _build_executive_summary(
        profile_name,
        profile,
        recommended_metrics,
        allocation_by_class,
        stress_results,
        rebalancing_trades,
        constraints_validation,
    )
    methodology = pd.DataFrame(
        [
            {"section": "Purpose", "description": "Educational workflow for IPS-driven ETF allocation and monitoring."},
            {"section": "Data", "description": f"Explicit {(provenance or {}).get('mode', 'unverified')} mode. See Data Provenance for per-asset sources and hashes; no silent source mixing."},
            {"section": "Optimization", "description": "SLSQP optimization with no shorting, IPS asset-class ranges, and max ETF weight limits."},
            {"section": "Risk", "description": "Return, volatility, Sharpe, Sortino, drawdown, VaR, CVaR, beta, tracking error, and information ratio."},
            {"section": "Factors", "description": "Fama-French regression is included as a diagnostic, but explanatory power is limited for a full multi-asset portfolio."},
            {"section": "Equity Sleeve Factors", "description": "Equity ETF weights are normalized and analyzed separately to better explain equity return drivers."},
            {"section": "Risk Contribution", "description": "ETF contribution to annualized portfolio volatility is calculated from the covariance matrix and target weights."},
            {"section": "Rolling Risk", "description": "Rolling 12-month return, volatility, Sharpe ratio, and drawdown are tracked against the benchmark."},
            {"section": "Stress Testing", "description": "Deterministic macro and market scenarios applied by ETF risk bucket."},
            {"section": "Rebalancing", "description": "Synthetic current weights are compared with targets using IPS drift thresholds."},
            {"section": "Charts", "description": "; ".join(str(path) for path in chart_paths.values())},
        ]
    )
    frontier_export = efficient_frontier.copy()
    if len(frontier_export) > 250:
        frontier_export = frontier_export.head(250)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _write_executive_summary_dashboard(writer.book, summary)
        if provenance:
            writer.book["Executive Summary"].merge_cells("A2:F2")
            writer.book["Executive Summary"].row_dimensions[2].height = 24
            mode_note = writer.book["Executive Summary"].cell(row=2, column=1)
            mode_note.value = f"{provenance['mode'].upper()} inputs | same-sample diagnostics | see Data Provenance"
            mode_note.font = Font(bold=True, size=10, color="9C0006")
        profiles.to_excel(writer, sheet_name="IPS Profile", index=False)
        allocation.to_excel(writer, sheet_name="Recommended Allocation", index=False)
        frontier_export.to_excel(writer, sheet_name="Efficient Frontier", index=False)
        risk_summary[
            [
                "profile_name",
                "portfolio_name",
                "annualized_return",
                "annualized_volatility",
                "sharpe_ratio",
                "max_drawdown",
                "tracking_error_vs_benchmark",
                "information_ratio",
            ]
        ].to_excel(writer, sheet_name="Portfolio Comparison", index=False)
        risk_summary.to_excel(writer, sheet_name="Risk Metrics", index=False)
        factor_summary.to_excel(writer, sheet_name="Factor Exposure", index=False)
        equity_factor_summary.to_excel(writer, sheet_name="Equity Sleeve Factors", index=False)
        risk_contribution.to_excel(writer, sheet_name="Risk Contribution", index=False)
        rolling_risk_metrics.to_excel(writer, sheet_name="Rolling Risk", index=False)
        stress_results.to_excel(writer, sheet_name="Stress Tests", index=False)
        rebalancing_trades.to_excel(writer, sheet_name="Rebalancing Trades", index=False)
        etf_rationale.to_excel(writer, sheet_name="ETF Rationale", index=False)
        constraints_validation.to_excel(writer, sheet_name="Constraint Validation", index=False)
        methodology.to_excel(writer, sheet_name="Methodology", index=False)
        if provenance:
            pd.DataFrame([{"asset": asset, **details} for asset, details in provenance["returns"]["assets"].items()]).to_excel(writer, sheet_name="Data Provenance", index=False)
            pd.DataFrame([{"section": "Returns", "interpretation": provenance["mode"],
                           "monthly_returns_sha256": provenance["returns"]["monthly_returns_sha256"],
                           "complete_months": provenance["returns"]["complete_months"]},
                          {"section": "Factors", **provenance["factors"]},
                          {"section": "Evaluation", "interpretation": provenance["evaluation"]}]).to_excel(writer, sheet_name="Source Summary", index=False)

    _format_workbook(output_path)
    return output_path


def generate_investment_memo(
    output_path: Path,
    profile: pd.Series,
    recommended_weights: pd.Series,
    risk_summary: pd.DataFrame,
    factor_summary: pd.DataFrame,
    equity_factor_summary: pd.DataFrame,
    risk_contribution: pd.DataFrame,
    rolling_risk_metrics: pd.DataFrame,
    etf_rationale: pd.DataFrame,
    constraints_validation: pd.DataFrame,
    stress_results: pd.DataFrame,
    rebalancing_trades: pd.DataFrame,
    provenance: dict | None = None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    profile_name = str(profile["profile_name"])
    metrics = risk_summary[
        (risk_summary["profile_name"] == profile_name)
        & (risk_summary["portfolio_name"] == "IPS Recommended Portfolio")
    ].iloc[0]
    allocation_by_class = class_allocation(recommended_weights)
    stress = stress_results[stress_results["profile_name"] == profile_name].drop_duplicates("scenario_name")
    stress = stress.sort_values("portfolio_scenario_return")
    trades = rebalancing_trades[rebalancing_trades["profile_name"] == profile_name]
    trade_actions = trades[trades["drift_breach_flag"]].copy()
    trade_lines = "\n".join(
        f"- {row['ticker']}: {row['trade_direction']} ${abs(row['trade_amount']):,.0f}; "
        f"drift {row['drift']:.2%}; target {row['target_weight']:.2%}; "
        f"estimated cost ${row['estimated_transaction_cost']:,.0f}"
        for _, row in trade_actions.iterrows()
    )
    if not trade_lines:
        trade_lines = "- No ETF drift exceeded the profile threshold in this demonstration run."
    worst_stress = stress.iloc[0]
    best_stress = stress.iloc[-1]
    relative_stress = stress.sort_values("relative_performance").iloc[0]
    class_text = ", ".join(
        f"{asset_class}: {allocation_by_class.get(asset_class, 0.0):.1%}" for asset_class in ASSET_CLASS_ORDER
    )
    full_r2 = factor_summary[
        (factor_summary["profile_name"] == profile_name) & (factor_summary["factor"] == "Alpha")
    ]["r_squared"].iloc[0]
    equity_r2 = equity_factor_summary[
        (equity_factor_summary["profile_name"] == profile_name) & (equity_factor_summary["factor"] == "Alpha")
    ]["r_squared"].iloc[0]
    top_risk = risk_contribution.sort_values("percent_risk_contribution", ascending=False).head(3)
    top_risk_text = ", ".join(
        f"{row['ticker']} ({row['percent_risk_contribution']:.1%})" for _, row in top_risk.iterrows()
    )
    risk_concentration_comment = (
        "Risk contribution is reasonably distributed across multiple ETFs."
        if top_risk["percent_risk_contribution"].sum() < 0.65
        else "Risk contribution is concentrated in the largest market-sensitive holdings."
    )
    latest_rolling = rolling_risk_metrics.sort_values("date").groupby("series").tail(1)
    latest_portfolio = latest_rolling[latest_rolling["series"] == "Recommended Portfolio"].iloc[0]
    latest_benchmark = latest_rolling[latest_rolling["series"] == "Benchmark"].iloc[0]
    constraints_passed = bool((constraints_validation["pass_fail"] == "PASS").all())
    constraint_sentence = (
        "The recommended allocation passed the IPS constraints validation."
        if constraints_passed
        else "One or more IPS constraint checks failed and should be reviewed before committee use."
    )
    sleeve_summary = etf_rationale.groupby("asset_class")["target_weight"].sum().to_dict()

    memo = f"""# Investment Committee Memo

## Executive Recommendation

This educational analysis recommends the **{profile_name}** IPS allocation as a model portfolio for committee review. The recommendation is framed as an educational investment workflow, not financial advice, not a live investment recommendation, and not production investment software.

The allocation is designed to remain within IPS ranges, preserve core benchmark anchors, limit high-yield and single-ETF concentration, and maintain practical diversification across growth, defensive, real asset, and liquidity exposures.

## IPS Constraints

The IPS objective is: {profile["return_objective"]}

- Maximum volatility target: {profile["max_volatility_target"]:.1%}
- Maximum drawdown tolerance: {profile["max_drawdown_tolerance"]:.1%}
- Maximum single ETF weight: {profile["max_single_etf_weight"]:.1%}
- Rebalancing threshold: {profile["rebalancing_threshold"]:.1%}
- Benchmark definition: {profile["benchmark_definition"]}

{constraint_sentence}

## Recommended Allocation

Recommended ETF allocation:

{format_weight_dict(recommended_weights)}

Asset-class allocation:

{class_text}

## Allocation Rationale by Sleeve

**Equity sleeve:** The equity allocation combines core U.S. equity exposure through SPY, a measured growth tilt through QQQ, small-cap exposure through IWM, and non-U.S. diversification through EFA and EEM. The international and emerging market allocation is capped to avoid excessive concentration in one non-U.S. risk source.

**Fixed income sleeve:** AGG provides the core bond ballast. TLT adds duration exposure that may help in equity-led growth shocks, while SHY dampens rate sensitivity. LQD and HYG add credit income, with HYG capped because high yield can behave like equity during spread widening.

**Alternatives sleeve:** VNQ and GLD provide real asset and crisis-hedge exposure. DBC remains available as a commodity sleeve where policy allows, but the Balanced Growth model does not force a commodity allocation when the optimizer and constraints do not require it.

**Cash sleeve:** SGOV serves as the cash/T-bill proxy. This sleeve supports liquidity, reduces forced selling risk, and provides dry powder for rebalancing.

Sleeve weights are approximately Equity {sleeve_summary.get("Equity", 0.0):.1%}, Fixed Income {sleeve_summary.get("Fixed Income", 0.0):.1%}, Alternatives {sleeve_summary.get("Alternatives", 0.0):.1%}, and Cash {sleeve_summary.get("Cash", 0.0):.1%}.

## Risk and Return Profile

- Annualized return: {metrics["annualized_return"]:.2%}
- Annualized volatility: {metrics["annualized_volatility"]:.2%}
- Sharpe ratio: {metrics["sharpe_ratio"]:.2f}
- Sortino ratio: {metrics["sortino_ratio"]:.2f}
- Max drawdown: {metrics["max_drawdown"]:.2%}
- Tracking error vs benchmark: {metrics["tracking_error_vs_benchmark"]:.2%}
- Information ratio: {metrics["information_ratio"]:.2f}
- 95% historical VaR: {metrics["historical_var_95"]:.2%}
- 95% historical CVaR: {metrics["historical_cvar_95"]:.2%}

Rolling monitoring shows the latest 12-month portfolio return at {latest_portfolio["rolling_12m_return"]:.2%}, volatility at {latest_portfolio["rolling_12m_volatility"]:.2%}, and rolling drawdown at {latest_portfolio["rolling_max_drawdown"]:.2%}. The latest benchmark readings are {latest_benchmark["rolling_12m_return"]:.2%}, {latest_benchmark["rolling_12m_volatility"]:.2%}, and {latest_benchmark["rolling_max_drawdown"]:.2%}, respectively. This monitoring view is intended to show how risk changes over time rather than relying only on full-period averages.

## Benchmark Comparison

The recommended portfolio is compared against the IPS benchmark definition. Its beta to the benchmark is {metrics["beta_vs_benchmark"]:.2f}, with correlation of {metrics["correlation_to_benchmark"]:.2f}. Tracking error of {metrics["tracking_error_vs_benchmark"]:.2%} reflects active allocation choices across growth equity, international equity, duration, credit, alternatives, and cash. The benchmark remains the reference point, but the recommended allocation deliberately broadens the sources of risk and return.

## Risk Contribution Summary

Risk contribution analysis estimates which ETFs contribute most to total portfolio volatility using the covariance matrix and target weights. The largest estimated contributors are {top_risk_text}. {risk_concentration_comment} This view is useful because a moderate weight can still contribute materially to risk if its volatility or covariance with the rest of the portfolio is high.

## Factor Exposure Interpretation

The factor regression shows limited explanatory power for the full multi-asset portfolio, which is expected because the portfolio includes bonds, credit, gold, cash, and alternatives. The full-portfolio regression R-squared is {full_r2:.2f}. These results should be treated as a diagnostic rather than a complete explanation of total portfolio returns.

## Equity-Sleeve Factor Analysis

The equity-sleeve regression normalizes only the equity ETF weights to 100% and runs the same Fama-French regression on that equity-only return stream. This is more appropriate for explaining equity return drivers because the regression is no longer diluted by bonds, cash, gold, and alternatives. The equity-sleeve regression R-squared is {equity_r2:.2f}. A future version would extend this approach with separate fixed-income attribution for duration, curve, and credit-spread exposures.

## Stress-Test Summary

The weakest modeled scenario is **{worst_stress["scenario_name"]}**, with an estimated portfolio return of {worst_stress["portfolio_scenario_return"]:.2%} versus benchmark return of {worst_stress["benchmark_scenario_return"]:.2%}. This is a realistic reminder that diversification can reduce, but not eliminate, broad market losses. The strongest modeled scenario is **{best_stress["scenario_name"]}**, with an estimated portfolio return of {best_stress["portfolio_scenario_return"]:.2%}. The least favorable relative scenario is **{relative_stress["scenario_name"]}**, where the portfolio trails the benchmark by {abs(relative_stress["relative_performance"]):.2%}.

## Rebalancing Actions

Synthetic current weights were generated to demonstrate policy monitoring. {len(trade_actions)} ETFs breached the IPS drift threshold. The resulting trades move breached positions back to target weights while leaving smaller drifts untouched.

{trade_lines}

Estimated transaction costs are ${trades["estimated_transaction_cost"].sum():,.0f} on a $100,000,000 illustrative portfolio.

## Key Risks

- The declared source mode is a sample assumption; synthetic results are not historical market performance and neither mode predicts future outcomes.
- Optimization is sensitive to expected returns, covariance estimates, and the selected risk-free-rate assumption.
- ETF proxies simplify implementation and do not capture manager selection, taxes, liquidity tiers, or mandate-specific restrictions.
- Stress tests are deterministic approximations and should be expanded for real committee use.
- Factor models are less informative for full multi-asset portfolios than for isolated equity sleeves.

## Limitations

This project is an educational portfolio analytics workflow. It is not a substitute for fiduciary review, due diligence, legal review, tax analysis, or live investment governance.

## Future Improvements

- Separate equity-sleeve and fixed-income-sleeve attribution.
- Transaction-lot-level rebalancing.
- Black-Litterman expected return inputs.
- Regime-aware optimization.
- Tax-aware rebalancing.
- Liquidity scoring.
- Manager and fund due diligence layer.
"""
    if provenance:
        memo = memo.replace("## Executive Recommendation", "## Data provenance\n\n"
                            f"Source mode: **{provenance['mode']}**. ETF source: {next(iter(provenance['returns']['assets'].values()))['source']}. "
                            f"Factor source: {provenance['factors']['source']}. "
                            "No market/synthetic asset substitution or zero return filling is applied. "
                            f"{provenance['evaluation']}. Inspect data_provenance.json and the workbook's Data Provenance sheet for every asset.\n\n"
                            "## Executive Recommendation")
    output_path.write_text(memo, encoding="utf-8")
    return output_path
