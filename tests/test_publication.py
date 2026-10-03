"""Check report consistency across real pipeline failures and saved readback."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

import main
from src import publication
from src.factor_model import generate_fallback_factors
from src.fetch_data import generate_fallback_monthly_returns


ROOT = Path(__file__).resolve().parents[1]


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2), encoding="utf-8")


def refresh_manifest_hash(root, relative):
    path = root / publication.MANIFEST_PATH
    manifest = read_json(path)
    digest = hashlib.sha256((root / relative).read_bytes()).hexdigest()
    manifest["files"][relative] = digest
    # Update optional grouped records too, so semantic tests do not stop at hashes.
    for key in ("inputs", "outputs", "input_sha256", "output_sha256"):
        group = manifest.get(key)
        if isinstance(group, dict) and relative in group:
            group[relative] = digest
    write_json(path, manifest)


def pack_bytes(root, include_status=False):
    names = [*publication.OWNED_FILES, publication.MANIFEST_PATH]
    if include_status:
        names.append(publication.STATUS_PATH)
    return {str(name): (root / name).read_bytes() for name in names if (root / name).is_file()}


@pytest.fixture(scope="module")
def baseline(tmp_path_factory):
    project = tmp_path_factory.mktemp("ips-complete-pack")
    main.run_pipeline(project, mode="synthetic", end="2026-05-31")
    publication.verify_pack(project)
    return project


@pytest.fixture
def project(baseline, tmp_path):
    target = tmp_path / "project"
    shutil.copytree(baseline, target)
    return target


def assert_failed_and_unchanged(project, before):
    assert pack_bytes(project) == before
    status = read_json(project / publication.STATUS_PATH)
    assert status["status"] == "FAILED"
    manifest = read_json(project / publication.MANIFEST_PATH)
    assert status["run_id"] != manifest["run_id"]
    with pytest.raises(publication.PublicationError):
        publication.verify_pack(project)
    assert not list(project.glob(".ips-stage-*"))


def test_completed_pack_matches_inputs_workbook_and_portable_chart_paths(baseline):
    manifest = publication.verify_pack(baseline)
    assert len(manifest["files"]) == 32
    assert read_json(baseline / publication.STATUS_PATH)["run_id"] == manifest["run_id"]
    for name, digest in manifest["files"].items():
        assert hashlib.sha256((baseline / name).read_bytes()).hexdigest() == digest
    provenance = read_json(baseline / "output/data_provenance.json")
    assert provenance["run_id"] == manifest["run_id"]
    assert provenance["returns"]["end_date"] == "2026-05-31"
    summary = pd.read_excel(baseline / "output/investment_committee_report.xlsx", sheet_name="Source Summary")
    returns = summary.loc[summary["section"] == "Returns"].iloc[0]
    assert returns["monthly_returns_sha256"] == provenance["returns"]["monthly_returns_sha256"]
    assert returns["end_date"] == "2026-05-31"
    publication_row = summary.loc[summary["section"] == "Publication"].iloc[0]
    assert publication_row["run_id"] == manifest["run_id"]
    assert publication_row["requested_end"] == "2026-05-31"
    for sheet, filename in (("Recommended Allocation", "recommended_allocation.csv"),
                            ("Risk Metrics", "portfolio_risk_summary.csv")):
        actual = pd.read_excel(baseline / "output/investment_committee_report.xlsx", sheet_name=sheet)
        expected = pd.read_csv(baseline / "output" / filename)
        pd.testing.assert_frame_equal(actual, expected, check_dtype=False, rtol=1e-10, atol=1e-12)
    methodology = pd.read_excel(baseline / "output/investment_committee_report.xlsx", sheet_name="Methodology")
    references = str(methodology.loc[methodology["section"] == "Charts", "description"].iloc[0])
    assert ".staging" not in references
    for reference in references.split("; "):
        normalized = reference.replace("\\", "/")
        assert normalized.startswith("charts/") and (baseline / normalized).is_file()


@pytest.mark.parametrize("provider", ["returns", "factors"])
def test_provider_failure_preserves_every_prior_artifact(monkeypatch, project, provider):
    before = pack_bytes(project)
    failure = ConnectionError(f"{provider} provider unavailable")
    name = "fetch_or_generate_monthly_returns" if provider == "returns" else "fetch_or_generate_factors"
    def fail(*args, **kwargs):
        raise failure
    def premature_analytics(*args, **kwargs):
        raise AssertionError("Analytics ran before both inputs were acquired")
    monkeypatch.setattr(main, name, fail)
    monkeypatch.setattr(main, "build_portfolio_candidates", premature_analytics)
    with pytest.raises(ConnectionError) as caught:
        main.run_pipeline(project, mode="synthetic", end="2026-04-30")
    assert caught.value is failure
    assert_failed_and_unchanged(project, before)


def test_late_workbook_format_failure_keeps_complete_previous_pack(monkeypatch, project, capsys):
    before = pack_bytes(project)
    failure = PermissionError("Workbook locked during final formatting")
    def fail(path):
        # ExcelWriter already saved a real candidate workbook inside staging.
        assert path.is_file()
        assert path.resolve() != (project / "output/investment_committee_report.xlsx").resolve()
        raise failure
    monkeypatch.setattr("src.generate_report._format_workbook", fail)
    with pytest.raises(PermissionError) as caught:
        main.run_pipeline(project, mode="synthetic", end="2026-04-30")
    assert caught.value is failure
    assert_failed_and_unchanged(project, before)
    assert "Pipeline complete" not in capsys.readouterr().out


def test_partial_chart_failure_preserves_complete_previous_pack(monkeypatch, project):
    before = pack_bytes(project)
    def fail(charts, *args, **kwargs):
        (charts / "recommended_allocation.png").write_bytes(b"partial candidate chart")
        raise RuntimeError("Chart generation interrupted")
    monkeypatch.setattr(main, "generate_charts", fail)
    with pytest.raises(RuntimeError, match="interrupted"):
        main.run_pipeline(project, mode="synthetic", end="2026-04-30")
    assert_failed_and_unchanged(project, before)


def test_publication_replace_failure_restores_already_replaced_files(monkeypatch, project):
    before = pack_bytes(project)
    target = project / "output/portfolio_risk_summary.csv"
    replace = Path.replace
    def fail_candidate(source, destination):
        if Path(destination).resolve() == target.resolve():
            raise PermissionError("Report table locked during publication")
        return replace(source, destination)
    monkeypatch.setattr(Path, "replace", fail_candidate)
    with pytest.raises(PermissionError, match="publication"):
        main.run_pipeline(project, mode="synthetic", end="2026-04-30")
    assert_failed_and_unchanged(project, before)


def test_failed_rollback_retains_original_pack_and_original_error(monkeypatch, project):
    before = pack_bytes(project)
    failed_target = project / "output/portfolio_risk_summary.csv"
    blocked_restore = project / "data/processed/monthly_returns.csv"
    original_replace, original_copy = Path.replace, shutil.copyfile
    failure = PermissionError("Original publication failure")
    def fail_candidate(source, destination):
        if Path(destination).resolve() == failed_target.resolve():
            raise failure
        return original_replace(source, destination)
    def fail_restore(source, destination, *args, **kwargs):
        if (any(part.startswith(".ips-recovery-") for part in Path(source).parts)
                and Path(destination).resolve() == blocked_restore.resolve()):
            raise PermissionError("Previous input is locked during rollback")
        return original_copy(source, destination, *args, **kwargs)
    monkeypatch.setattr(Path, "replace", fail_candidate)
    monkeypatch.setattr(shutil, "copyfile", fail_restore)
    with pytest.raises(PermissionError) as caught:
        main.run_pipeline(project, mode="synthetic", end="2026-04-30")
    assert caught.value is failure
    status = read_json(project / publication.STATUS_PATH)
    assert status["status"] == "FAILED"
    recovery = Path(status["recovery_path"])
    assert recovery.resolve().parent == project.resolve()
    assert recovery.name.startswith(".ips-recovery-")
    for name, content in before.items():
        assert (recovery / name).read_bytes() == content
    assert any("preserved" in note for note in failure.__notes__)
    assert not list(project.glob(".ips-stage-*"))
    with pytest.raises(publication.PublicationError):
        publication.verify_pack(project)


def test_failure_writing_success_marker_rolls_back_report_and_manifest(monkeypatch, project):
    before = pack_bytes(project)
    status_path = project / publication.STATUS_PATH
    replace = Path.replace
    def fail_success(source, destination):
        if (Path(destination).resolve() == status_path.resolve()
                and read_json(source).get("status") == "SUCCESS"):
            raise PermissionError("Cannot commit SUCCESS marker")
        return replace(source, destination)
    monkeypatch.setattr(Path, "replace", fail_success)
    with pytest.raises(PermissionError, match="SUCCESS"):
        main.run_pipeline(project, mode="synthetic", end="2026-04-30")
    assert_failed_and_unchanged(project, before)


def test_market_pipeline_binds_raw_prices_and_both_provider_sources(monkeypatch, tmp_path):
    # These are synthetic provider doubles, not market performance evidence.
    returns = generate_fallback_monthly_returns(end_date=pd.Timestamp("2026-05-31"))
    dates = pd.DatetimeIndex([returns.index[0] - pd.offsets.MonthEnd(1), *returns.index])
    prices = pd.concat([pd.DataFrame([[100.0] * len(returns.columns)], columns=returns.columns),
                        (1 + returns).cumprod().mul(100).reset_index(drop=True)], ignore_index=True)
    prices.columns = returns.columns
    prices.index = dates
    factors = generate_fallback_factors(returns.index) * 100
    factors.index = factors.index.to_period("M")
    calls = []
    def market_prices(**kwargs):
        calls.append("prices")
        return prices
    def market_factors(*args, **kwargs):
        calls.append("factors")
        return {0: factors}
    monkeypatch.setattr("src.fetch_data.fetch_daily_prices", market_prices)
    monkeypatch.setattr("pandas_datareader.data.DataReader", market_factors)
    project = tmp_path / "market-pack"
    main.run_pipeline(project, mode="market", end="2026-05-31")
    manifest = publication.verify_pack(project)
    assert calls == ["prices", "factors"]
    assert len(manifest["files"]) == 33
    assert manifest["mode"] == "market"
    assert "data/raw/daily_etf_prices.csv" in manifest["inputs"]
    provenance = read_json(project / "output/data_provenance.json")
    assert provenance["returns"]["daily_prices_sha256"] == manifest["inputs"]["data/raw/daily_etf_prices.csv"]
    assert all(not item["synthetic"] for item in provenance["returns"]["assets"].values())


@pytest.mark.parametrize("end", ["NaT", "2026-02-30", "2026-05-31T12:00:00"])
def test_invalid_calendar_options_do_not_start_or_mutate_a_run(monkeypatch, project, end):
    before = pack_bytes(project, include_status=True)
    def forbidden(*args, **kwargs):
        raise AssertionError("Invalid date reached input acquisition")
    monkeypatch.setattr(main, "fetch_or_generate_monthly_returns", forbidden)
    with pytest.raises(ValueError):
        main.run_pipeline(project, mode="synthetic", end=end)
    assert pack_bytes(project, include_status=True) == before
    publication.verify_pack(project)


def test_successful_retry_removes_stale_market_file_and_keeps_user_files(monkeypatch, project):
    failure = ConnectionError("Factors unavailable")
    provider = main.fetch_or_generate_factors
    monkeypatch.setattr(main, "fetch_or_generate_factors", lambda *args, **kwargs: (_ for _ in ()).throw(failure))
    with pytest.raises(ConnectionError):
        main.run_pipeline(project, mode="synthetic", end="2026-04-30")
    monkeypatch.setattr(main, "fetch_or_generate_factors", provider)
    stale = project / "data/raw/daily_etf_prices.csv"
    stale.write_text("old market data", encoding="utf-8")
    notes = [project / "charts/personal_notes.txt", project / "output/custom_report.txt"]
    for note in notes:
        note.write_text("preserve user material", encoding="utf-8")
    main.run_pipeline(project, mode="synthetic", end="2026-04-30")
    manifest = publication.verify_pack(project)
    assert not stale.exists()
    assert all(note.read_text(encoding="utf-8") == "preserve user material" for note in notes)
    assert read_json(project / publication.STATUS_PATH)["status"] == "SUCCESS"
    assert read_json(project / "output/data_provenance.json")["returns"]["end_date"] == "2026-04-30"
    assert manifest["run_id"] == read_json(project / publication.STATUS_PATH)["run_id"]


@pytest.mark.parametrize("name", ["output/investment_committee_report.xlsx", "output/investment_memo.md",
                                "data/processed/monthly_returns.csv", "charts/recommended_allocation.png"])
def test_changed_owned_files_are_rejected(project, name):
    with (project / name).open("ab") as handle:
        handle.write(b"altered after publication")
    with pytest.raises(publication.PublicationError):
        publication.verify_pack(project)


@pytest.mark.parametrize("name", [publication.STATUS_PATH, publication.MANIFEST_PATH,
                                "data/processed/fama_french_factors.csv"])
def test_missing_controls_or_inputs_are_rejected(project, name):
    (project / name).unlink()
    with pytest.raises(publication.PublicationError):
        publication.verify_pack(project)


@pytest.mark.parametrize("status", ["RUNNING", "FAILED"])
def test_unfinished_or_failed_status_cannot_validate_old_report(project, status):
    path = project / publication.STATUS_PATH
    saved = read_json(path)
    saved["status"] = status
    write_json(path, saved)
    with pytest.raises(publication.PublicationError):
        publication.verify_pack(project)


def test_other_attempt_id_cannot_validate_pack(project):
    path = project / publication.STATUS_PATH
    saved = read_json(path)
    saved["run_id"] = "different-attempt"
    write_json(path, saved)
    with pytest.raises(publication.PublicationError):
        publication.verify_pack(project)


@pytest.mark.parametrize("change", ["mode", "date", "allocation", "workbook", "profile"])
def test_rehashed_semantic_disagreement_is_rejected(project, change):
    if change in {"mode", "date"}:
        relative = "output/data_provenance.json"
        data = read_json(project / relative)
        if change == "mode":
            data["mode"] = "market"
        else:
            data["returns"]["end_date"] = "2026-04-30"
        write_json(project / relative, data)
    elif change == "allocation":
        relative = "output/recommended_allocation.csv"
        data = pd.read_csv(project / relative)
        data.loc[0, "weight"] += 0.01
        data.to_csv(project / relative, index=False)
    elif change == "profile":
        relative = "data/raw/ips_profiles.csv"
        data = pd.read_csv(project / relative)
        data.loc[0, "max_single_etf_weight"] += 0.01
        data.to_csv(project / relative, index=False)
    else:
        relative = "output/investment_committee_report.xlsx"
        workbook = load_workbook(project / relative)
        sheet = workbook["Source Summary"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        section_column = headers["section"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, section_column).value == "Returns":
                sheet.cell(row, headers["monthly_returns_sha256"], "wrong-source-hash")
        workbook.save(project / relative)
        workbook.close()
    refresh_manifest_hash(project, relative)
    with pytest.raises(publication.PublicationError):
        publication.verify_pack(project)


def test_overlapping_run_and_verifier_stop_without_changing_current_pack(project):
    before = pack_bytes(project, include_status=True)
    with publication.output_lock(project):
        with pytest.raises(publication.PublicationError):
            main.run_pipeline(project, mode="synthetic", end="2026-04-30")
        with pytest.raises(publication.PublicationError):
            publication.verify_pack(project)
    assert pack_bytes(project, include_status=True) == before
    publication.verify_pack(project)


def test_cli_verification_and_failure_exit_status(project):
    command = [sys.executable, str(ROOT / "main.py"), "--project-dir", str(project), "--verify-only"]
    completed = subprocess.run(command, cwd=ROOT, capture_output=True, text=True)
    assert completed.returncode == 0, completed.stderr
    path = project / publication.STATUS_PATH
    saved = read_json(path)
    saved["status"] = "FAILED"
    write_json(path, saved)
    before = pack_bytes(project, include_status=True)
    rejected = subprocess.run(command, cwd=ROOT, capture_output=True, text=True)
    assert rejected.returncode != 0
    assert pack_bytes(project, include_status=True) == before
