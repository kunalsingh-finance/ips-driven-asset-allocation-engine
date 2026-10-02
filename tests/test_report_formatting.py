from openpyxl import Workbook, load_workbook

from src.generate_report import _format_workbook


def test_market_source_summary_without_synthetic_seed_formats(tmp_path):
    path = tmp_path / "market.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Source Summary"
    sheet.append(["mode", "seed"])
    sheet.append(["market", None])
    workbook.save(path)

    _format_workbook(path)

    saved = load_workbook(path)
    assert saved["Source Summary"]["B2"].value is None
    assert saved["Source Summary"].column_dimensions["B"].width == 6
    saved.close()


def test_header_only_metric_sheet_does_not_create_reversed_format_range(tmp_path):
    path = tmp_path / "empty.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Risk Summary"
    sheet.append(["annualized_return", "max_drawdown", "pass_fail"])
    workbook.save(path)

    _format_workbook(path)

    saved = load_workbook(path)
    assert saved["Risk Summary"].max_row == 1
    assert len(saved["Risk Summary"].conditional_formatting) == 0
    saved.close()
